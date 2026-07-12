#!/usr/bin/env python3
"""
DTIClaw Excel/CSV Writer
Generate a styled .xlsx (or .csv) from JSON or CSV input.

Input JSON shape (one or many sheets):
  {"sheets": [{"name": "Sheet1", "rows": [["A","B"], [1,2]]}]}
  or a flat list of rows: [["A","B"],[1,2]]

Usage:
  python3 xlsx_writer.py --input data.json --output report.xlsx --title "Laporan"
  python3 xlsx_writer.py --input data.csv  --output report.xlsx
  python3 xlsx_writer.py --input data.json --output report.csv   # CSV (first sheet)
"""
import argparse
import csv
import json
import os
import sys


def load_rows(path):
    """Return a list of sheets: [{name, rows}]."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8") as f:
            return [{"name": "Sheet1", "rows": [r for r in csv.reader(f)]}]
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return [{"name": "Sheet1", "rows": data}]
    if isinstance(data, dict) and "sheets" in data:
        return data["sheets"]
    if isinstance(data, dict) and "rows" in data:
        return [{"name": data.get("name", "Sheet1"), "rows": data["rows"]}]
    raise SystemExit("ERROR: unsupported JSON shape (need list, {rows}, or {sheets}).")


def write_csv(sheets, output):
    rows = sheets[0]["rows"]
    with open(output, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    print(f"✅ CSV saved: {output} ({len(rows)} rows)")


def write_xlsx(sheets, output, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1A3C6E")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet in sheets:
        ws = wb.create_sheet(title=(sheet.get("name") or "Sheet")[:31])
        rows = sheet.get("rows", [])
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = border
                if r_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")
        # Auto width
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(width + 3, 60)
        ws.freeze_panes = "A2"

    if title:
        wb.properties.title = title
    os.makedirs(os.path.dirname(os.path.abspath(output)), exist_ok=True)
    wb.save(output)
    total = sum(len(s.get("rows", [])) for s in sheets)
    print(f"✅ Excel saved: {output} ({len(sheets)} sheet, {total} rows)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="DTIClaw Excel/CSV writer")
    ap.add_argument("--input", required=True, help="JSON or CSV input file")
    ap.add_argument("--output", required=True, help="Output .xlsx or .csv")
    ap.add_argument("--title", default="", help="Document title (xlsx metadata)")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        sys.exit(f"ERROR: input not found: {args.input}")
    sheets = load_rows(args.input)
    if args.output.lower().endswith(".csv"):
        write_csv(sheets, args.output)
    else:
        write_xlsx(sheets, args.output, args.title)
